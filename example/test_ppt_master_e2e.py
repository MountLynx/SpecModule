# example/test_ppt_master_e2e.py
"""Mock 全链 E2E：4 页册（≤6 无早门）、notes 关（最少条件段）、零 LLM。

ScriptedMock 按 prompt 内容分流：计划 prompt 是唯一含收据 schema 串
"roster_ids" 的 prompt（_PLAN_CORE 收据提示原样透传；页/修复/备注 prompt
与全部 prompt_extra 素材均无此串——Task 5/9 实测核实）→ 计划收据 JSON；
其余 → fixture SVG。全链真跑 vendor checker/finalize/export——工程复刻
的验收底线。

计划修订实测锚（Task 1/5/6 + 本任务预检结论）：
- fixture SVG 在 ``fixtures/svg_output/page_p01.svg``（Task 1 实测路径）。
- spec_lock 收据用 fixture 全文：checker final 门要求 typography 的
  title/body 行与 ``pptx_structure.mode``（本任务预检实测：草稿
  "typography: sans" 记法触发 blocking "Master export requires
  spec_lock.md typography title and body rows"）；fixture 锁已过
  Task 1 四件套。plan_validate 还要求锁文本含 "palette" 锚点词
  （Task 6），fixture 锁无此词——标题行补锚点词，不动 checker 实测
  通过的节/行结构。
- 4 个相同 SVG 预检实测不触发 checker blocking（重复类检查只针对
  defs id 与 layout manifest 指纹，不跨页比较内容）——无需按页参数化。
- design_spec 收据不能只写 plan_validate 的 `## pNN` 页块：design_spec.md
  在盘时 checker communication_trace 还要求 `## IX. Content Outline` +
  `#### Slide NN` 块（每块含 Audience move 与 Relationships 行）——
  本任务首跑实测（scope-only blocking → repair 空 转 2 轮 → 上限
  infrastructure Failure loud abort，修复环语义符合 Task 6 设计）。
"""

from __future__ import annotations

import asyncio
import json
from pathlib import Path

from example.ppt_master import workspace
from example.ppt_master.module import run_generate
from llm.client import LLMResponse
from llm.mock import MockLLMClient

FIXTURES = Path(__file__).parent / "ppt_master" / "fixtures"
_FIXTURE_SVG = FIXTURES / "svg_output" / "page_p01.svg"
_FIXTURE_LOCK = FIXTURES / "spec_lock.md"

_SVG = _FIXTURE_SVG.read_text(encoding="utf-8")
# fixture 锁全文（Task 1 实测过 checker final 门）+ 标题行补 plan_validate
# 的 palette 锚点词（Task 6）；节/行结构与实测通过的 fixture 完全一致
_LOCK = _FIXTURE_LOCK.read_text(encoding="utf-8").replace(
    "# Execution Lock",
    "# Execution Lock (palette anchored)",
)
assert "palette" in _LOCK   # 绊线：fixture 标题若变，此处即失败（非远处 plan_validate）
# `#### Slide NN` 块喂 checker communication_trace（design_spec.md 在盘时的
# 门要求）；`## pNN` 块喂 plan_validate——Audience move 双写是两校验器重叠要求。
_SPEC_MD = (
    "# spec\n"
    "## IX. Content Outline\n\n"
    "#### Slide 01 - p01 封面\n"
    "- **Audience move**: 建立第一印象\n"
    "- **Relationships**: 标题与要点总分\n\n"
    "#### Slide 02 - p02 方法\n"
    "- **Audience move**: 交代路径\n"
    "- **Relationships**: 步骤序列\n\n"
    "#### Slide 03 - p03 结果\n"
    "- **Audience move**: 呈现证据\n"
    "- **Relationships**: 证据并列\n\n"
    "#### Slide 04 - p04 结语\n"
    "- **Audience move**: 给出结论\n"
    "- **Relationships**: 收束呼应\n\n"
    "## p01 封面\n"
    "Audience move: 建立第一印象\n\n"
    "## p02 方法\n"
    "Audience move: 交代路径\n\n"
    "## p03 结果\n"
    "Audience move: 呈现证据\n\n"
    "## p04 结语\n"
    "Audience move: 给出结论\n"
)


class ScriptedMock(MockLLMClient):
    """按 prompt 分流的脚本客户端：计划收据 JSON / fixture SVG。"""

    async def complete(self, **kw):
        prompt = kw.get("prompt", "")
        if "roster_ids" in prompt:  # 仅计划 prompt 含此收据 schema 串
            return LLMResponse(content=json.dumps({
                "status": "ok",
                "roster_ids": ["p01", "p02", "p03", "p04"],
                "design_spec_md": _SPEC_MD,
                "spec_lock_md": _LOCK,
                "image_rows": [], "icon_pool": [], "notes_enabled": False,
            }))
        return LLMResponse(content=_SVG)


def _spec(output_dir: Path) -> dict:
    return {
        "project": "e2e",
        # 源文件不存在：ingest 对缺失路径容错（digest 记"契约驱动"，Task 6）
        "source": {"kind": "files",
                   "paths": [str(output_dir.parent / "sources_absent.md")]},
        "production": {"speaker_notes": False},
        "output": {"dir": str(output_dir)},
        "roster": [
            {"id": "p01", "title": "封面", "role": "cover"},
            {"id": "p02", "title": "方法"},
            {"id": "p03", "title": "结果"},
            {"id": "p04", "title": "结语", "role": "closing"},
        ],
    }


def test_full_pipeline_mock(tmp_path, monkeypatch):
    """零 LLM 全链：翻译 → 规划收据 → 4 页 fixture SVG → 真 checker final 门
    → finalize → svg_to_pptx（--no-notes）→ Report ok + 4 页 pptx 在盘。"""
    monkeypatch.setattr(workspace, "_ENVELOPE_DIR", tmp_path)
    out_dir = tmp_path / "deck"
    firings = asyncio.run(run_generate(
        _spec(out_dir), llm_client=ScriptedMock(), persist=False))
    by_node = {f.node: f.output for f in firings}

    report = by_node.get("Report", {})
    assert report.get("status") == "ok", (
        f"Report 收据: {report}（fired 节点: {sorted(by_node)}）")
    # 门一次过（fixture SVG 已验过 final 门）：verdict 干净、零修复轮
    assert by_node["FinalVerdict"]["ok"] is True, by_node.get("FinalVerdict")
    assert "Repair" not in by_node, by_node.get("Repair")
    assert by_node["FinalVerdict"]["returncode"] == 0

    pptx = list((out_dir / "exports").glob("*.pptx"))
    assert pptx, "exports/ 应有 pptx"

    from pptx import Presentation
    assert len(Presentation(str(pptx[0])).slides) == 4


def test_cli_mock_smoke(tmp_path, monkeypatch, capsys):
    """CLI 冒烟：真 entry 发现 + 命令面 + --mock 通道全链跑通。

    CLI --mock 硬编码 ``MockLLMClient()``（cli._build_llm_client，无注入
    参数/环境面）：其固定收据 {"result": ...} 无 roster_ids/design_spec，
    plan_validate 必然 infrastructure Failure——这是 ppt_master 契约面
    "无隐式兜底" 的应有行为，非缺陷。且 tickflow 对 infrastructure
    Failure 是返回而非抛出，CLI rc 恒 0——降断言（只看 rc）是空断言。
    故经 monkeypatch 替换 CLI 名字空间的 MockLLMClient 符号为
    ScriptedMock：--mock 通道（免 key 假客户端）语义不变，仅脚本体可产
    合法收据。全链断言主责在 test_full_pipeline_mock，本测试断言 CLI 面
    通畅（rc=0 + pptx 落盘）。
    """
    import module_harness.cli.cli as cli_mod

    monkeypatch.setattr(cli_mod, "MockLLMClient", ScriptedMock)
    monkeypatch.setattr(workspace, "_ENVELOPE_DIR", tmp_path)  # 信封不落共享 %TEMP%
    monkeypatch.chdir(tmp_path)
    from module_harness.cli import main
    spec = _spec(tmp_path / "cli_deck")
    rc = main([
        "run", "--module", "ppt_master", "--mock",
        "--modules-dir", str(Path(__file__).parent / "modules"),
        "--run-id", "ppt_master_cli_smoke",
        "--spec", json.dumps(spec, ensure_ascii=False),
    ])
    assert rc == 0, capsys.readouterr().err
    pptx = list((tmp_path / "cli_deck" / "exports").glob("*.pptx"))
    assert pptx, "exports/ 应有 pptx"
    from pptx import Presentation
    assert len(Presentation(str(pptx[0])).slides) == 4
